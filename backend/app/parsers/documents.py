import csv
import io

import pymupdf  # PyMuPDF
from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation

from .base import Page, ParseError
from .ocr import ocr_bytes, ocr_image

MIN_CHARS_PER_PDF_PAGE = 30  # below this we assume a scanned page and OCR it


def decode_text(data: bytes) -> str:
    # utf-8-sig strips a BOM; utf-16 covers files saved by some Windows tools
    for enc in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def parse_pdf(data: bytes) -> list[Page]:
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception:
        raise ParseError("Could not open PDF (corrupt or password-protected).")
    pages = []
    for i, page in enumerate(doc, start=1):
        text = page.get_text().strip()
        if len(text) < MIN_CHARS_PER_PDF_PAGE:  # scanned page -> OCR fallback
            pix = page.get_pixmap(dpi=200)
            img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            text = ocr_image(img).strip()
        if text:
            pages.append(Page(i, text))
    return pages


def parse_docx(data: bytes) -> list[Page]:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return [Page(1, "\n".join(parts))] if parts else []


def parse_pptx(data: bytes) -> list[Page]:
    prs = Presentation(io.BytesIO(data))
    pages = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = [s.text_frame.text for s in slide.shapes
                 if s.has_text_frame and s.text_frame.text.strip()]
        if texts:
            pages.append(Page(i, "\n".join(texts)))
    return pages


def parse_xlsx(data: bytes) -> list[Page]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    pages = []
    for i, ws in enumerate(wb.worksheets, start=1):
        rows = [" | ".join("" if c is None else str(c) for c in row)
                for row in ws.iter_rows(values_only=True)]
        text = "\n".join(r for r in rows if r.strip(" |"))
        if text:
            pages.append(Page(i, f"[Sheet: {ws.title}]\n{text}"))
    return pages


def parse_csv(data: bytes) -> list[Page]:
    rows = [" | ".join(r) for r in csv.reader(io.StringIO(decode_text(data)))]
    return [Page(1, "\n".join(rows))] if rows else []


def parse_html(data: bytes) -> list[Page]:
    soup = BeautifulSoup(decode_text(data), "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    text = soup.get_text("\n", strip=True)
    return [Page(1, text)] if text else []


def parse_text(data: bytes) -> list[Page]:
    text = decode_text(data).strip()
    return [Page(1, text)] if text else []


def parse_image(data: bytes) -> list[Page]:
    text = ocr_bytes(data).strip()
    return [Page(1, text)] if text else []
